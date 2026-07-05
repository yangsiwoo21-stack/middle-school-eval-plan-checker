from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from checker import (
    CODE_RE,
    Document,
    ReviewFinding,
    assessment_overview_items,
    assessment_ratio_table_rows,
    build_excel_relation_audit,
    extract_basic_score_rows,
    normalize_area_name,
    performance_area_names_from_ratio,
    performance_area_percentages_from_ratio,
    performance_area_scores_from_ratio,
    performance_detail_area_names,
    sector_text,
    split_combined_subject_document,
    split_performance_blocks,
)


@dataclass
class AuditRow:
    sheet: str
    values: dict[str, object]


class ExcelReviewBuilder:
    """Builds the v0.3 Excel-based inspection workbook.

    The workbook is intentionally used as an intermediate audit table, not only
    as a final export. Each row keeps the extracted value, the sector source,
    and a Pass/Fail/확인 필요 status so false positives can be traced quickly.
    """

    SUMMARY_HEADERS = ["파일", "학년", "과목", "섹터", "판정", "설명"]
    MONTHLY_HEADERS = ["파일", "학년", "과목", "월", "주", "성취기준", "평가방법/연계", "원문"]
    OVERVIEW_HEADERS = ["파일", "학년", "과목", "영역명", "반영비율", "영역만점", "성취기준", "평가시기", "판정", "설명"]
    DETAIL_HEADERS = ["파일", "학년", "과목", "영역명", "영역만점", "기본점수", "기본점수비율", "성취기준", "판정", "설명"]
    COMPARE_HEADERS = ["파일", "학년", "과목", "비교구역", "비교항목", "기준값", "비교값", "판정", "검토 의견"]
    FINDING_HEADERS = ["파일", "학년", "과목", "심각도", "분류", "판정", "근거", "수정 제안"]

    def build(self, paths: list[Path], findings: list[ReviewFinding], output_path: Path) -> Path:
        workbook = Workbook()
        default = workbook.active
        default.title = "추출요약"
        sheets = {
            "추출요약": default,
            "월별계획": workbook.create_sheet("월별계획"),
            "평가반영비율": workbook.create_sheet("평가반영비율"),
            "수행평가세부기준": workbook.create_sheet("수행평가세부기준"),
            "구역간비교": workbook.create_sheet("구역간비교"),
            "진단결과": workbook.create_sheet("진단결과"),
        }
        headers = {
            "추출요약": self.SUMMARY_HEADERS,
            "월별계획": self.MONTHLY_HEADERS,
            "평가반영비율": self.OVERVIEW_HEADERS,
            "수행평가세부기준": self.DETAIL_HEADERS,
            "구역간비교": self.COMPARE_HEADERS,
            "진단결과": self.FINDING_HEADERS,
        }
        for name, header in headers.items():
            self._write_header(sheets[name], header)

        documents = self._load_documents(paths)
        for document in documents:
            self._append_document_rows(sheets, document)

        for finding in findings:
            self._append_dict(
                sheets["진단결과"],
                self.FINDING_HEADERS,
                {
                    "파일": finding.file_path.name,
                    "학년": finding.grade or "",
                    "과목": finding.subject or "",
                    "심각도": finding.severity,
                    "분류": finding.topic,
                    "판정": "Fail" if finding.severity in {"상", "중"} else "확인 필요",
                    "근거": finding.anchor_text,
                    "수정 제안": finding.memo_text.replace("\n", " / "),
                },
            )

        for sheet in sheets.values():
            self._format_sheet(sheet)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path

    def _load_documents(self, paths: list[Path]) -> list[Document]:
        documents: list[Document] = []
        for path in paths:
            if path.is_dir():
                files = sorted([*path.rglob("*.hwpx"), *path.rglob("*.md")])
            else:
                files = [path]
            for file_path in files:
                if file_path.suffix.lower() not in {".hwpx", ".md"}:
                    continue
                try:
                    document = Document.from_path(file_path)
                except Exception:
                    continue
                documents.extend(split_combined_subject_document(document))
        return documents

    def _append_document_rows(self, sheets: dict[str, object], document: Document) -> None:
        self._append_summary_rows(sheets["추출요약"], document)
        self._append_monthly_rows(sheets["월별계획"], document)
        self._append_overview_rows(sheets["평가반영비율"], document)
        self._append_detail_rows(sheets["수행평가세부기준"], document)
        self._append_compare_rows(sheets["구역간비교"], document)

    def _append_summary_rows(self, sheet, document: Document) -> None:
        required = [
            ("document_info", "문서 기본 정보"),
            ("monthly_plan", "교수·학습 운영 계획"),
            ("assessment_overview", "평가의 종류와 반영비율"),
            ("performance_detail", "수행평가 세부기준"),
            ("absence", "미응시자 및 학적변동자 처리"),
        ]
        for sector_id, label in required:
            exists = sector_id in document.sectors
            self._append_dict(
                sheet,
                self.SUMMARY_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "섹터": label,
                    "판정": "Pass" if exists else "확인 필요",
                    "설명": "표준 섹터 인식" if exists else "섹터 제목 변형 또는 추출 실패 가능",
                },
            )

    def _append_monthly_rows(self, sheet, document: Document) -> None:
        text = sector_text(document, "monthly_plan")
        if not text:
            return
        rows = extract_monthly_plan_rows(text)
        if not rows:
            self._append_dict(
                sheet,
                self.MONTHLY_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "월": "",
                    "주": "",
                    "성취기준": "",
                    "평가방법/연계": "",
                    "원문": text[:500],
                },
            )
            return
        for row in rows:
            row.update({"파일": document.path.name, "학년": document.grade or "", "과목": document.subject or ""})
            self._append_dict(sheet, self.MONTHLY_HEADERS, row)

    def _append_overview_rows(self, sheet, document: Document) -> None:
        items = assessment_overview_items(document)
        if not items:
            self._append_dict(
                sheet,
                self.OVERVIEW_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "영역명": "",
                    "판정": "확인 필요",
                    "설명": "평가의 종류와 반영비율 표에서 수행평가 영역명을 추출하지 못함",
                },
            )
            return
        for item in items:
            ratio = item.ratio if item.ratio is not None else ""
            score = item.score if item.score is not None else ""
            status = "Pass"
            note = "추출 성공"
            if isinstance(ratio, (int, float)) and ratio > 30:
                status = "Fail"
                note = "수행평가 한 영역 30% 초과"
            self._append_dict(
                sheet,
                self.OVERVIEW_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "영역명": item.name,
                    "반영비율": ratio,
                    "영역만점": score,
                    "성취기준": ", ".join(item.achievement_codes),
                    "평가시기": item.period,
                    "판정": status,
                    "설명": note,
                },
            )

    def _append_detail_rows(self, sheet, document: Document) -> None:
        section = sector_text(document, "performance_detail", "6. 수행평가", "7.")
        if not section:
            self._append_dict(
                sheet,
                self.DETAIL_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "판정": "확인 필요",
                    "설명": "수행평가 세부기준 섹터를 추출하지 못함",
                },
            )
            return

        basic_rows = extract_basic_score_rows(section)
        if basic_rows:
            for full_score, basic_score, context, _anchor in basic_rows:
                percent = round(basic_score / full_score * 100, 1) if full_score else ""
                status = "Pass" if isinstance(percent, float) and 20 <= percent <= 40 else "Fail"
                self._append_dict(
                    sheet,
                    self.DETAIL_HEADERS,
                    {
                        "파일": document.path.name,
                        "학년": document.grade or "",
                        "과목": document.subject or "",
                        "영역명": extract_area_name(context),
                        "영역만점": full_score,
                        "기본점수": basic_score,
                        "기본점수비율": percent,
                        "성취기준": ", ".join(sorted(set(CODE_RE.findall(context)))),
                        "판정": status,
                        "설명": "기본점수 20~40% 범위 확인" if status == "Pass" else "기본점수 범위 오류",
                    },
                )
            return

        for block in split_performance_blocks(section):
            self._append_dict(
                sheet,
                self.DETAIL_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "영역명": extract_area_name(block),
                    "영역만점": extract_area_score(block),
                    "성취기준": ", ".join(sorted(set(CODE_RE.findall(block)))),
                    "판정": "확인 필요",
                    "설명": "기본점수 행을 명확히 추출하지 못함",
                },
            )

    def _append_compare_rows(self, sheet, document: Document) -> None:
        for relation in build_excel_relation_audit(document):
            self._append_dict(
                sheet,
                self.COMPARE_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "비교구역": relation.compare_area,
                    "비교항목": relation.check_item,
                    "기준값": relation.left_value,
                    "비교값": relation.right_value,
                    "판정": relation.status,
                    "검토 의견": relation.opinion,
                },
            )

        overview_names = performance_area_names_from_ratio(document)
        detail_names = performance_detail_area_names(document)
        normalized_details = {normalize_area_name(name): name for name in detail_names}
        for name in overview_names:
            key = normalize_area_name(name)
            found = key in normalized_details or any(key and (key in detail or detail in key) for detail in normalized_details)
            self._append_dict(
                sheet,
                self.COMPARE_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "비교구역": "assessment_overview ↔ performance_detail",
                    "비교항목": "수행평가 영역명",
                    "기준값": name,
                    "비교값": normalized_details.get(key, ""),
                    "판정": "Pass" if found else "Fail",
                    "검토 의견": "4번 평가표와 6번 세부기준 영역명 일치" if found else "6번 세부기준에서 같은 영역명을 확인하지 못함",
                },
            )

        monthly_text = sector_text(document, "monthly_plan")
        normalized_monthly = normalize_area_name(monthly_text)
        for name in overview_names:
            key = normalize_area_name(name)
            found = bool(key and key in normalized_monthly)
            self._append_dict(
                sheet,
                self.COMPARE_HEADERS,
                {
                    "파일": document.path.name,
                    "학년": document.grade or "",
                    "과목": document.subject or "",
                    "비교구역": "monthly_plan ↔ assessment_overview",
                    "비교항목": "수행평가 연계",
                    "기준값": name,
                    "비교값": "월별계획 내 확인" if found else "",
                    "판정": "Pass" if found else "확인 필요",
                    "검토 의견": "월별 교수학습 운영계획에 수행평가명이 확인됨" if found else "수행평가명이 다르게 표현되었거나 월별계획에 누락되었을 수 있음",
                },
            )

    def _append_dict(self, sheet, headers: list[str], row: dict[str, object]) -> None:
        sheet.append([row.get(header, "") for header in headers])

    def _write_header(self, sheet, headers: list[str]) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="335C67")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _format_sheet(self, sheet) -> None:
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.value == "Pass":
                    cell.fill = PatternFill("solid", fgColor="D9EAD3")
                elif cell.value == "Fail":
                    cell.fill = PatternFill("solid", fgColor="F4CCCC")
                elif cell.value == "확인 필요":
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
        for column_cells in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            width = min(max(length + 2, 10), 48)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def extract_monthly_plan_rows(text: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    current_month = ""
    for line in [clean_line(line) for line in text.splitlines() if clean_line(line)]:
        month_match = re.fullmatch(r"([0-9]{1,2})\s*월", line)
        if month_match:
            current_month = month_match.group(1)
            continue
        if "수행평가" not in line and not CODE_RE.search(line):
            continue
        week_match = re.search(r"([1-5])\s*주", line)
        rows.append(
            {
                "월": current_month,
                "주": week_match.group(1) if week_match else "",
                "성취기준": ", ".join(sorted(set(CODE_RE.findall(line)))),
                "평가방법/연계": "수행평가" if "수행평가" in line else "",
                "원문": line[:500],
            }
        )
    return rows[:300]


def extract_area_name(text: str) -> str:
    patterns = [
        r"평가영역명\s*\n?\s*([^\n|()]{2,80})",
        r"(?:[가-하]\.|[0-9]+\))\s*([^\n()]{2,80})(?:\(|\n)",
        r"([^\n|()]{2,80})\((\d{1,3})점\)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return clean_line(match.group(1))
    return ""


def extract_area_score(text: str) -> object:
    match = re.search(r"\((\d{1,3})점\)", text)
    return int(match.group(1)) if match else ""


def extract_period_near_name(text: str, name: str) -> str:
    if not text or not name:
        return ""
    idx = normalize_area_name(text).find(normalize_area_name(name))
    if idx < 0:
        return ""
    window = text[max(0, idx - 300): idx + 800]
    match = re.search(r"([0-9]{1,2})\s*월\s*([1-5])\s*주", window)
    return f"{match.group(1)}월 {match.group(2)}주" if match else ""


def clean_line(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()
